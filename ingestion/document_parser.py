"""
FinCo AI - Document Parser

Responsible for:
- Parsing PDF, DOCX, TXT, and Markdown documents
- Extracting readable text
- Preserving page-level information where possible
- Cleaning extracted text
- Preparing document content for the RAG pipeline
- Returning structured parsing results

Supported formats:
    - .pdf
    - .docx
    - .txt
    - .md
    - .markdown

Example:

    from app.ingestion.document_parser import document_parser

    result = document_parser.parse(
        "data/raw/financial_report.pdf"
    )

    print(result.text)
    print(result.metadata)
"""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, BinaryIO, Dict, List, Optional

logger = logging.getLogger(__name__)


# ============================================================================
# Exceptions
# ============================================================================


class DocumentParserError(Exception):
    """Base exception for document parsing errors."""


class DocumentFileNotFoundError(DocumentParserError):
    """Raised when a document file cannot be found."""


class UnsupportedDocumentTypeError(DocumentParserError):
    """Raised when the document format is not supported."""


class DocumentSizeLimitError(DocumentParserError):
    """Raised when a document exceeds the configured size limit."""


class DocumentExtractionError(DocumentParserError):
    """Raised when text extraction fails."""


# ============================================================================
# Configuration
# ============================================================================


@dataclass(frozen=True)
class DocumentParserConfig:
    """
    Configuration for document parsing.
    """

    supported_extensions: tuple[str, ...] = (
        ".pdf",
        ".docx",
        ".txt",
        ".md",
        ".markdown",
    )

    max_file_size_mb: int = 50

    min_text_length: int = 1

    default_text_encoding: str = "utf-8"

    remove_repeated_whitespace: bool = True

    preserve_page_breaks: bool = True

    include_empty_pages: bool = False

    extract_docx_tables: bool = True

    extract_pdf_metadata: bool = True


# ============================================================================
# Result Models
# ============================================================================


@dataclass
class DocumentPage:
    """
    Represents one logical document page.

    For TXT, Markdown, and DOCX files, page_number may represent
    a logical section rather than a physical page.
    """

    page_number: int

    text: str

    metadata: Dict[str, Any] = field(
        default_factory=dict
    )

    @property
    def character_count(self) -> int:
        return len(self.text)

    @property
    def word_count(self) -> int:
        return len(self.text.split())

    def to_dict(self) -> Dict[str, Any]:
        return {
            "page_number": self.page_number,
            "text": self.text,
            "character_count": self.character_count,
            "word_count": self.word_count,
            "metadata": self.metadata,
        }


@dataclass
class DocumentMetadata:
    """
    Metadata associated with a parsed document.
    """

    source_name: Optional[str] = None

    source_path: Optional[str] = None

    extension: Optional[str] = None

    content_type: Optional[str] = None

    file_size_bytes: int = 0

    page_count: int = 0

    character_count: int = 0

    word_count: int = 0

    title: Optional[str] = None

    author: Optional[str] = None

    subject: Optional[str] = None

    creation_date: Optional[str] = None

    modified_date: Optional[str] = None

    extra: Dict[str, Any] = field(
        default_factory=dict
    )

    def to_dict(self) -> Dict[str, Any]:
        return {
            "source_name": self.source_name,
            "source_path": self.source_path,
            "extension": self.extension,
            "content_type": self.content_type,
            "file_size_bytes": self.file_size_bytes,
            "page_count": self.page_count,
            "character_count": self.character_count,
            "word_count": self.word_count,
            "title": self.title,
            "author": self.author,
            "subject": self.subject,
            "creation_date": self.creation_date,
            "modified_date": self.modified_date,
            "extra": self.extra,
        }


@dataclass
class DocumentParseResult:
    """
    Result returned by DocumentParser.parse().
    """

    text: str

    pages: List[DocumentPage]

    metadata: DocumentMetadata

    success: bool = True

    warnings: List[str] = field(
        default_factory=list
    )

    errors: List[str] = field(
        default_factory=list
    )

    @property
    def page_count(self) -> int:
        return len(self.pages)

    @property
    def character_count(self) -> int:
        return len(self.text)

    @property
    def word_count(self) -> int:
        return len(self.text.split())

    def to_dict(self) -> Dict[str, Any]:
        return {
            "text": self.text,
            "pages": [
                page.to_dict()
                for page in self.pages
            ],
            "metadata": self.metadata.to_dict(),
            "success": self.success,
            "warnings": self.warnings,
            "errors": self.errors,
        }


# ============================================================================
# Document Parser
# ============================================================================


class DocumentParser:
    """
    Parser for financial documents used by FinCo AI.

    This class extracts text only. OCR for scanned PDFs should be
    implemented as a separate service or fallback pipeline.
    """

    def __init__(
        self,
        config: DocumentParserConfig = DocumentParserConfig(),
    ) -> None:
        self.config = config

    # ---------------------------------------------------------------------
    # Public API
    # ---------------------------------------------------------------------

    def parse(
        self,
        source: str | Path | bytes | bytearray | BinaryIO,
        source_name: Optional[str] = None,
    ) -> DocumentParseResult:
        """
        Parse a supported document source.

        Supported source types:
            - File path
            - pathlib.Path
            - bytes
            - bytearray
            - BinaryIO
        """

        resolved_name = (
            source_name
            or self._get_source_name(source)
            or "unknown_document"
        )

        raw_bytes = self._read_source(source)

        self._validate_file_size(
            raw_bytes
        )

        extension = self._resolve_extension(
            source=source,
            source_name=resolved_name,
        )

        metadata = DocumentMetadata(
            source_name=resolved_name,
            source_path=(
                str(source)
                if isinstance(source, (str, Path))
                else None
            ),
            extension=extension,
            content_type=self._get_content_type(extension),
            file_size_bytes=len(raw_bytes),
        )

        try:
            if extension == ".pdf":
                pages, extra_metadata = self._parse_pdf(
                    raw_bytes
                )

            elif extension == ".docx":
                pages, extra_metadata = self._parse_docx(
                    raw_bytes
                )

            elif extension in {
                ".txt",
                ".md",
                ".markdown",
            }:
                pages, extra_metadata = self._parse_text(
                    raw_bytes
                )

            else:
                raise UnsupportedDocumentTypeError(
                    f"Unsupported document extension: {extension}"
                )

        except DocumentParserError:
            raise

        except Exception as exc:
            logger.exception(
                "Document extraction failed: %s",
                resolved_name,
            )

            raise DocumentExtractionError(
                f"Failed to extract document text: {exc}"
            ) from exc

        pages = self._clean_pages(pages)

        text = self._combine_pages(pages)

        if len(text.strip()) < self.config.min_text_length:
            raise DocumentExtractionError(
                "Document does not contain enough extractable text. "
                "The file may be scanned or image-only."
            )

        metadata.page_count = len(pages)
        metadata.character_count = len(text)
        metadata.word_count = len(text.split())
        metadata.extra.update(extra_metadata)

        warnings: List[str] = []

        if extension == ".pdf" and not text.strip():
            warnings.append(
                "PDF contains no extractable text. OCR may be required."
            )

        return DocumentParseResult(
            text=text,
            pages=pages,
            metadata=metadata,
            success=True,
            warnings=warnings,
        )

    # ---------------------------------------------------------------------
    # Source handling
    # ---------------------------------------------------------------------

    def _get_source_name(
        self,
        source: Any,
    ) -> Optional[str]:
        if isinstance(source, (str, Path)):
            return Path(source).name

        return getattr(source, "name", None)

    def _read_source(
        self,
        source: str | Path | bytes | bytearray | BinaryIO,
    ) -> bytes:
        if isinstance(source, Path):
            if not source.exists():
                raise DocumentFileNotFoundError(
                    f"Document file not found: {source}"
                )

            if not source.is_file():
                raise DocumentFileNotFoundError(
                    f"Document source is not a file: {source}"
                )

            return source.read_bytes()

        if isinstance(source, str):
            path = Path(source)

            if not path.exists():
                raise DocumentFileNotFoundError(
                    f"Document file not found: {source}"
                )

            if not path.is_file():
                raise DocumentFileNotFoundError(
                    f"Document source is not a file: {source}"
                )

            return path.read_bytes()

        if isinstance(source, bytes):
            return source

        if isinstance(source, bytearray):
            return bytes(source)

        if hasattr(source, "read"):
            content = source.read()

            if isinstance(content, str):
                return content.encode(
                    self.config.default_text_encoding
                )

            return content

        raise TypeError(
            "Unsupported document source type. "
            "Use a path, bytes, bytearray, or binary file object."
        )

    def _validate_file_size(
        self,
        raw_bytes: bytes,
    ) -> None:
        max_bytes = (
            self.config.max_file_size_mb
            * 1024
            * 1024
        )

        if len(raw_bytes) > max_bytes:
            raise DocumentSizeLimitError(
                f"Document size is {len(raw_bytes)} bytes. "
                f"Maximum allowed size is {max_bytes} bytes."
            )

    def _resolve_extension(
        self,
        source: Any,
        source_name: str,
    ) -> str:
        if isinstance(source, (str, Path)):
            extension = Path(source).suffix.lower()

        else:
            extension = Path(source_name).suffix.lower()

        if not extension:
            raise UnsupportedDocumentTypeError(
                "Unable to determine document type. "
                "Provide a source_name with a file extension."
            )

        if extension not in self.config.supported_extensions:
            raise UnsupportedDocumentTypeError(
                f"Unsupported document type: {extension}. "
                f"Supported types: "
                f"{self.config.supported_extensions}"
            )

        return extension

    def _get_content_type(
        self,
        extension: str,
    ) -> str:
        content_types = {
            ".pdf": "application/pdf",
            ".docx": (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            ".txt": "text/plain",
            ".md": "text/markdown",
            ".markdown": "text/markdown",
        }

        return content_types.get(
            extension,
            "application/octet-stream",
        )

    # ---------------------------------------------------------------------
    # PDF parsing
    # ---------------------------------------------------------------------

    def _parse_pdf(
        self,
        raw_bytes: bytes,
    ) -> tuple[List[DocumentPage], Dict[str, Any]]:
        """
        Extracts text from PDF pages using pypdf.
        """

        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise DocumentExtractionError(
                "pypdf is required to parse PDF files. "
                "Install it with: pip install pypdf"
            ) from exc

        try:
            reader = PdfReader(
                io.BytesIO(raw_bytes)
            )

            pages: List[DocumentPage] = []

            for index, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""

                if (
                    not page_text.strip()
                    and not self.config.include_empty_pages
                ):
                    continue

                pages.append(
                    DocumentPage(
                        page_number=index + 1,
                        text=page_text,
                        metadata={
                            "source_type": "pdf",
                        },
                    )
                )

            extra_metadata: Dict[str, Any] = {}

            if self.config.extract_pdf_metadata:
                pdf_metadata = reader.metadata

                if pdf_metadata:
                    extra_metadata["pdf_metadata"] = {
                        str(key): str(value)
                        for key, value in pdf_metadata.items()
                        if value is not None
                    }

            return pages, extra_metadata

        except Exception as exc:
            raise DocumentExtractionError(
                f"Unable to parse PDF: {exc}"
            ) from exc

    # ---------------------------------------------------------------------
    # DOCX parsing
    # ---------------------------------------------------------------------

    def _parse_docx(
        self,
        raw_bytes: bytes,
    ) -> tuple[List[DocumentPage], Dict[str, Any]]:
        """
        Extracts paragraphs and tables from DOCX files using python-docx.
        """

        try:
            from docx import Document
        except ImportError as exc:
            raise DocumentExtractionError(
                "python-docx is required to parse DOCX files. "
                "Install it with: pip install python-docx"
            ) from exc

        try:
            document = Document(
                io.BytesIO(raw_bytes)
            )

            content: List[str] = []

            for paragraph in document.paragraphs:
                text = paragraph.text.strip()

                if text:
                    content.append(text)

            if self.config.extract_docx_tables:
                for table in document.tables:
                    table_text = self._extract_docx_table(
                        table
                    )

                    if table_text:
                        content.append(table_text)

            combined_text = "\n\n".join(content)

            pages = [
                DocumentPage(
                    page_number=1,
                    text=combined_text,
                    metadata={
                        "source_type": "docx",
                    },
                )
            ]

            core_properties = document.core_properties

            extra_metadata = {
                "docx_metadata": {
                    "title": core_properties.title,
                    "author": core_properties.author,
                    "subject": core_properties.subject,
                    "created": (
                        str(core_properties.created)
                        if core_properties.created
                        else None
                    ),
                    "modified": (
                        str(core_properties.modified)
                        if core_properties.modified
                        else None
                    ),
                }
            }

            return pages, extra_metadata

        except Exception as exc:
            raise DocumentExtractionError(
                f"Unable to parse DOCX: {exc}"
            ) from exc

    def _extract_docx_table(
        self,
        table: Any,
    ) -> str:
        rows: List[str] = []

        for row in table.rows:
            cells = []

            for cell in row.cells:
                cell_text = " ".join(
                    paragraph.text.strip()
                    for paragraph in cell.paragraphs
                    if paragraph.text.strip()
                )

                cells.append(cell_text)

            rows.append(
                " | ".join(cells)
            )

        return "\n".join(rows)

    # ---------------------------------------------------------------------
    # TXT and Markdown parsing
    # ---------------------------------------------------------------------

    def _parse_text(
        self,
        raw_bytes: bytes,
    ) -> tuple[List[DocumentPage], Dict[str, Any]]:
        """
        Parses TXT and Markdown files.
        """

        text = self._decode_text(
            raw_bytes
        )

        pages = [
            DocumentPage(
                page_number=1,
                text=text,
                metadata={
                    "source_type": "text",
                },
            )
        ]

        return pages, {}

    def _decode_text(
        self,
        raw_bytes: bytes,
    ) -> str:
        encodings = (
            "utf-8-sig",
            "utf-8",
            "cp1252",
            "latin-1",
        )

        for encoding in encodings:
            try:
                return raw_bytes.decode(
                    encoding
                )

            except UnicodeDecodeError:
                continue

        raise DocumentExtractionError(
            "Unable to decode text document."
        )

    # ---------------------------------------------------------------------
    # Text cleaning
    # ---------------------------------------------------------------------

    def _clean_pages(
        self,
        pages: List[DocumentPage],
    ) -> List[DocumentPage]:
        cleaned_pages: List[DocumentPage] = []

        for page in pages:
            cleaned_text = self.clean_text(
                page.text
            )

            if (
                not cleaned_text
                and not self.config.include_empty_pages
            ):
                continue

            cleaned_pages.append(
                DocumentPage(
                    page_number=page.page_number,
                    text=cleaned_text,
                    metadata=page.metadata,
                )
            )

        return cleaned_pages

    def clean_text(
        self,
        text: str,
    ) -> str:
        """
        Cleans extracted text without destroying paragraph structure.
        """

        if not text:
            return ""

        text = text.replace(
            "\x00",
            "",
        )

        text = text.replace(
            "\r\n",
            "\n",
        )

        text = text.replace(
            "\r",
            "\n",
        )

        # Remove trailing spaces from each line.
        text = "\n".join(
            line.rstrip()
            for line in text.splitlines()
        )

        if self.config.remove_repeated_whitespace:
            text = re.sub(
                r"[ \t]+",
                " ",
                text,
            )

            text = re.sub(
                r"\n{3,}",
                "\n\n",
                text,
            )

        return text.strip()

    def _combine_pages(
        self,
        pages: List[DocumentPage],
    ) -> str:
        if not pages:
            return ""

        page_separator = (
            "\n\n--- PAGE BREAK ---\n\n"
            if self.config.preserve_page_breaks
            else "\n\n"
        )

        return page_separator.join(
            page.text
            for page in pages
            if page.text.strip()
        )


# ============================================================================
# Singleton
# ============================================================================


document_parser = DocumentParser()